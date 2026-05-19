"""
Output Formatter — แปลง LLM output เป็น file format ต่างๆ
1 task → 1 file
"""

import os
import json
import re
from datetime import date
from typing import Optional


def safe_format(template: str, ctx: dict) -> str:
    """Replace {key} placeholders only for keys present in ctx; leave others untouched."""
    def replacer(m):
        key = m.group(1)
        return str(ctx[key]) if key in ctx else m.group(0)
    return re.sub(r'\{(\w+)\}', replacer, template)


def save_task_output(
    content: str,
    output_file: str,
    output_format: str,
    output_dir: str,
) -> str:
    """
    บันทึก content ลงไฟล์ตาม format ที่กำหนด
    คืน absolute path ของไฟล์ที่บันทึก

    Formats:
      excel      → openpyxl .xlsx (JSON→sheets or markdown table)
      word       → python-docx .docx (markdown→Word)
      mermaid    → raw .mmd (strip ```mermaid fences)
      code_multi → parse === FILE: path === blocks, write each file, manifest .md
      others     → write raw text
    """
    os.makedirs(output_dir, exist_ok=True)
    full_path = os.path.join(output_dir, output_file)
    os.makedirs(os.path.dirname(full_path), exist_ok=True)

    if output_format == "excel":
        _save_excel(content, full_path)

    elif output_format == "word":
        _save_docx(content, full_path)

    elif output_format == "mermaid":
        # .mmd files: save raw mermaid syntax without markdown fences
        mmd_content = _extract_mermaid_raw(content)
        with open(full_path, "w", encoding="utf-8") as f:
            f.write(mmd_content)

    elif output_format == "code_multi":
        # Save raw content as manifest, then extract individual code files
        with open(full_path, "w", encoding="utf-8") as f:
            f.write(content)
        _save_multi_file_code(content, output_dir)

    else:
        # markdown, html, sql, yaml, dockerfile, json — write raw
        with open(full_path, "w", encoding="utf-8") as f:
            f.write(content)

    return full_path


def _extract_mermaid_raw(content: str) -> str:
    """Extract pure mermaid syntax, stripping ```mermaid fences if present."""
    m = re.search(r"```mermaid\s*(.*?)\s*```", content, re.DOTALL)
    if m:
        return m.group(1).strip() + "\n"
    # Already raw mermaid — return as-is
    return content.strip() + "\n"


def _save_multi_file_code(content: str, output_dir: str):
    """
    Parse LLM output that uses === FILE: path/to/file.ext === markers.
    Each section is saved as an actual file under output_dir.

    Expected format:
        === FILE: src/main.py ===
        [file content]
        === FILE: src/utils.py ===
        [file content]
    """
    # Split on the marker pattern
    pattern = re.compile(r'^=== FILE: (.+?) ===$', re.MULTILINE)
    parts = pattern.split(content)

    # parts[0] = text before first marker (preamble, ignored)
    # parts[1::2] = file paths
    # parts[2::2] = file contents
    file_paths = parts[1::2]
    file_contents = parts[2::2]

    saved = []
    for fpath, fcontent in zip(file_paths, file_contents):
        fpath = fpath.strip()
        if not fpath:
            continue
        # Sanitize: prevent path traversal
        safe_path = os.path.normpath(fpath).lstrip("/").lstrip("\\")
        if ".." in safe_path:
            continue
        full = os.path.join(output_dir, safe_path)
        os.makedirs(os.path.dirname(full), exist_ok=True)
        # Strip leading/trailing blank lines from content
        cleaned = fcontent.strip("\n").rstrip()
        with open(full, "w", encoding="utf-8") as f:
            f.write(cleaned + "\n")
        saved.append(safe_path)

    return saved


def _save_excel(content: str, path: str):
    """
    content เป็น JSON ที่ LLM produce:
    {
      "sheets": [
        {
          "name": "RACI Matrix",
          "headers": ["Task", "CEO", "PM", "BA", "SA"],
          "rows": [["Task A", "R", "A", "C", "I"], ...]
        }
      ]
    }
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        with open(path.replace(".xlsx", ".md"), "w", encoding="utf-8") as f:
            f.write(f"# {os.path.basename(path)}\n\n```\n{content}\n```")
        return

    # แยก JSON จาก content
    data = _extract_json(content)
    if not data or "sheets" not in data:
        # ถ้า JSON ไม่ได้ — ลองแปลง markdown table
        data = _markdown_table_to_json(content)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # ลบ default sheet

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    alt_fill    = PatternFill("solid", fgColor="D6E4F7")
    border      = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )

    for sheet_def in data.get("sheets", [{"name": "Sheet1", "headers": [], "rows": []}]):
        ws = wb.create_sheet(title=sheet_def.get("name", "Sheet1")[:31])
        headers = sheet_def.get("headers", [])
        rows    = sheet_def.get("rows", [])

        # Header row
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = border

        # Data rows
        for r_idx, row in enumerate(rows, start=2):
            fill = alt_fill if r_idx % 2 == 0 else None
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border    = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if fill:
                    cell.fill = fill

        # Auto-fit columns (approximation)
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        ws.row_dimensions[1].height = 25

    wb.save(path)


def _save_docx(content: str, path: str):
    """
    Convert markdown content → python-docx Word document.
    Handles: # headings, **bold**, tables, - lists, 1. numbered lists, code blocks.
    Falls back to plain .md if python-docx unavailable.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        # Fallback: save as markdown
        fallback = path.rsplit(".", 1)[0] + ".md"
        with open(fallback, "w", encoding="utf-8") as f:
            f.write(content)
        return

    doc = Document()

    # Document styling
    style = doc.styles["Normal"]
    style.font.name = "Sarabun"
    style.font.size = Pt(11)

    def _set_heading_style(para, level: int):
        colors = {1: "1F4E79", 2: "2E74B5", 3: "2E74B5", 4: "595959"}
        run = para.runs[0] if para.runs else para.add_run()
        run.font.color.rgb = RGBColor.from_string(colors.get(level, "000000"))

    def _add_formatted_run(para, text: str):
        """Add text with inline bold (**text**) and italic (*text*) support."""
        # Split on bold markers
        segments = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
        for seg in segments:
            if seg.startswith("**") and seg.endswith("**"):
                run = para.add_run(seg[2:-2])
                run.bold = True
            elif seg.startswith("*") and seg.endswith("*"):
                run = para.add_run(seg[1:-1])
                run.italic = True
            elif seg.startswith("`") and seg.endswith("`"):
                run = para.add_run(seg[1:-1])
                run.font.name = "Courier New"
                run.font.size = Pt(10)
            elif seg:
                para.add_run(seg)

    def _add_table(table_lines: list):
        rows = [l for l in table_lines if not re.match(r'^\|[\s\-:|]+\|', l)]
        if not rows:
            return
        cells_per_row = [
            [c.strip() for c in r.strip().strip("|").split("|")]
            for r in rows
        ]
        if not cells_per_row:
            return
        ncols = max(len(r) for r in cells_per_row)
        table = doc.add_table(rows=len(cells_per_row), cols=ncols)
        table.style = "Table Grid"
        for r_idx, row_data in enumerate(cells_per_row):
            for c_idx in range(ncols):
                cell = table.cell(r_idx, c_idx)
                val = row_data[c_idx] if c_idx < len(row_data) else ""
                # Bold header row
                if r_idx == 0:
                    run = cell.paragraphs[0].add_run(val)
                    run.bold = True
                    # Blue background for header
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "D6E4F7")
                    tcPr.append(shd)
                else:
                    _add_formatted_run(cell.paragraphs[0], val)
        doc.add_paragraph()  # spacing after table

    lines = content.split("\n")
    i = 0
    in_code_block = False
    code_lines: list = []

    while i < len(lines):
        line = lines[i]

        # Code blocks
        if line.strip().startswith("```"):
            if in_code_block:
                # End code block — add as formatted paragraph
                in_code_block = False
                if code_lines:
                    para = doc.add_paragraph()
                    run = para.add_run("\n".join(code_lines))
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)
                code_lines = []
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Headings
        heading_m = re.match(r'^(#{1,4})\s+(.+)', line)
        if heading_m:
            level = len(heading_m.group(1))
            text = heading_m.group(2).strip()
            # Strip inline markdown from heading
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            para = doc.add_heading(text, level=level)
            _set_heading_style(para, level)
            i += 1
            continue

        # Tables — collect consecutive | lines
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            _add_table(table_lines)
            continue

        # Unordered list
        if re.match(r'^[\s]*[-*]\s+', line):
            text = re.sub(r'^[\s]*[-*]\s+', '', line)
            para = doc.add_paragraph(style="List Bullet")
            _add_formatted_run(para, text)
            i += 1
            continue

        # Numbered list
        if re.match(r'^[\s]*\d+\.\s+', line):
            text = re.sub(r'^[\s]*\d+\.\s+', '', line)
            para = doc.add_paragraph(style="List Number")
            _add_formatted_run(para, text)
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^[\s]*---+[\s]*$', line) or re.match(r'^[\s]*===+[\s]*$', line):
            i += 1
            continue

        # Empty line — add spacing
        if not line.strip():
            i += 1
            continue

        # Normal paragraph
        para = doc.add_paragraph()
        _add_formatted_run(para, line)
        i += 1

    doc.save(path)


def _extract_json(text: str) -> Optional[dict]:
    """พยายามแยก JSON จาก text"""
    # หา ```json ... ```
    m = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL)
    if m:
        try:
            return json.loads(m.group(1))
        except Exception:
            pass
    # หา { ... } โดยตรง
    start = text.find("{")
    end   = text.rfind("}") + 1
    if start >= 0 and end > start:
        try:
            return json.loads(text[start:end])
        except Exception:
            pass
    return None


def _markdown_table_to_json(text: str) -> dict:
    """แปลง Markdown table เป็น JSON sheets"""
    sheets = []
    current_headers: list = []
    current_rows: list = []
    sheet_name = "Sheet1"

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            if current_headers:
                sheets.append({"name": sheet_name, "headers": current_headers, "rows": current_rows})
            sheet_name = stripped.lstrip("#").strip()
            current_headers = []
            current_rows = []
        elif "|" in stripped and not stripped.startswith("|---"):
            parts = [p.strip() for p in stripped.strip("|").split("|")]
            if not current_headers:
                current_headers = parts
            else:
                current_rows.append(parts)

    if current_headers:
        sheets.append({"name": sheet_name, "headers": current_headers, "rows": current_rows})

    return {"sheets": sheets} if sheets else {
        "sheets": [{"name": "Data", "headers": ["Content"], "rows": [[text]]}]
    }


def extract_mermaid(content: str) -> str:
    """
    ถ้า content มี ```mermaid block ดึงออกมา
    ถ้าไม่มี ส่งคืน content เดิม (LLM อาจ output mermaid ตรงๆ)
    """
    m = re.search(r"```mermaid\s*(.*?)\s*```", content, re.DOTALL)
    if m:
        return f"```mermaid\n{m.group(1).strip()}\n```"
    # ถ้า content เริ่มด้วย keyword mermaid ให้ wrap
    triggers = ("graph ", "flowchart ", "sequenceDiagram", "erDiagram",
                "gantt", "classDiagram", "stateDiagram", "C4Context", "C4Container")
    if any(content.strip().startswith(t) for t in triggers):
        return f"```mermaid\n{content.strip()}\n```"
    return content


def wrap_mermaid_in_md(content: str, title: str = "") -> str:
    """Wrap mermaid code ใน markdown ที่อ่านง่าย"""
    mermaid_block = extract_mermaid(content)
    if title:
        return f"# {title}\n\n{mermaid_block}\n"
    return mermaid_block + "\n"
