"""
Attachment intake helper for CEO project creation.

Parses Discord message attachments into text suitable for project requirements.

Supported types:
  .txt / .md  — UTF-8 text (latin-1 fallback)
  .csv        — decoded and rendered as pipe-delimited rows
  .xlsx       — parsed with openpyxl (must be installed); sheet + row data
  .xls        — not supported (legacy format); user gets a conversion hint

Usage:
  combined, warnings = await gather_intake(ctx.message, inline_text)

  combined  — merged requirements string (inline text + attachment content)
  warnings  — list of user-facing warning strings to post to Discord
"""

import csv
import io
import os
from typing import NamedTuple

MAX_SINGLE_BYTES    = 4 * 1024 * 1024  # 4 MB raw-file size guard (before parse)
MAX_ATTACHMENT_CHARS = 15_000           # per-attachment character limit after parsing
MAX_TOTAL_CHARS      = 25_000           # combined inline + all attachments

SUPPORTED_TEXT_EXTS  = {".txt", ".md"}
SUPPORTED_CSV_EXTS   = {".csv"}
SUPPORTED_EXCEL_EXTS = {".xlsx"}         # .xls excluded — openpyxl does not support it
SUPPORTED_EXTS = SUPPORTED_TEXT_EXTS | SUPPORTED_CSV_EXTS | SUPPORTED_EXCEL_EXTS


class AttachmentResult(NamedTuple):
    filename:  str
    text:      str   # parsed text content (empty string on error)
    error:     str   # non-empty when parsing failed; safe to show to user
    truncated: bool  # True when content was cut at MAX_ATTACHMENT_CHARS


# ─── Per-format parsers ───────────────────────────────────────────────────────

def _ext(filename: str) -> str:
    return os.path.splitext(filename.lower())[1]


def _parse_text_bytes(data: bytes, filename: str) -> AttachmentResult:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="replace")

    truncated = len(text) > MAX_ATTACHMENT_CHARS
    if truncated:
        text = text[:MAX_ATTACHMENT_CHARS]
    return AttachmentResult(filename=filename, text=text, error="", truncated=truncated)


def _parse_csv_bytes(data: bytes, filename: str) -> AttachmentResult:
    try:
        raw = data.decode("utf-8")
    except UnicodeDecodeError:
        raw = data.decode("latin-1", errors="replace")

    rows = []
    for row in csv.reader(io.StringIO(raw)):
        rows.append(" | ".join(cell.strip() for cell in row))
    text = "\n".join(rows)

    truncated = len(text) > MAX_ATTACHMENT_CHARS
    if truncated:
        text = text[:MAX_ATTACHMENT_CHARS]
    return AttachmentResult(filename=filename, text=text, error="", truncated=truncated)


def _parse_xlsx_bytes(data: bytes, filename: str) -> AttachmentResult:
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        return AttachmentResult(
            filename=filename, text="",
            error="openpyxl is not installed — Excel support requires: pip install openpyxl",
            truncated=False,
        )

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"### Sheet: {sheet_name}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    except Exception as exc:
        return AttachmentResult(
            filename=filename, text="",
            error=f"Excel parse error: {exc}",
            truncated=False,
        )

    truncated = len(text) > MAX_ATTACHMENT_CHARS
    if truncated:
        text = text[:MAX_ATTACHMENT_CHARS]
    return AttachmentResult(filename=filename, text=text, error="", truncated=truncated)


# ─── Public API ───────────────────────────────────────────────────────────────

def parse_bytes(data: bytes, filename: str) -> AttachmentResult:
    """
    Parse raw attachment bytes by file extension.
    Safe to call synchronously; used in tests and by gather_intake.
    """
    ext = _ext(filename)
    if ext in SUPPORTED_TEXT_EXTS:
        return _parse_text_bytes(data, filename)
    if ext in SUPPORTED_CSV_EXTS:
        return _parse_csv_bytes(data, filename)
    if ext == ".xlsx":
        return _parse_xlsx_bytes(data, filename)
    if ext == ".xls":
        return AttachmentResult(
            filename=filename, text="",
            error=".xls (legacy Excel) is not supported. Please save as .xlsx and re-upload.",
            truncated=False,
        )
    return AttachmentResult(
        filename=filename, text="",
        error=(
            f"Unsupported file type '{ext}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTS))}"
        ),
        truncated=False,
    )


async def gather_intake(message, inline_text: str = "") -> tuple[str, list[str]]:
    """
    Gather requirements text from a Discord message.

    Combines inline_text with parsed attachments into a single string.
    Each attachment section is prefixed with ## Attachment: <filename>.

    Returns:
        combined  — merged text, ready for _start_project_from_text()
        warnings  — user-visible warning strings (truncation, errors, size)
    """
    parts: list[str] = []
    warnings: list[str] = []

    if inline_text.strip():
        parts.append(inline_text.strip())

    attachments = getattr(message, "attachments", []) or []
    for att in attachments:
        size = getattr(att, "size", 0)
        fname = getattr(att, "filename", "attachment")

        if size > MAX_SINGLE_BYTES:
            warnings.append(
                f"⚠️ `{fname}` ({size // 1024:,} KB) exceeds the "
                f"{MAX_SINGLE_BYTES // (1024 * 1024)} MB limit — skipped."
            )
            continue

        data = await att.read()
        result = parse_bytes(data, fname)

        if result.error:
            warnings.append(f"⚠️ `{fname}`: {result.error}")
            continue

        if result.truncated:
            warnings.append(
                f"⚠️ `{fname}` was truncated to {MAX_ATTACHMENT_CHARS:,} chars."
            )

        parts.append(f"## Attachment: {fname}\n{result.text}")

    combined = "\n\n".join(parts)

    if len(combined) > MAX_TOTAL_CHARS:
        combined = combined[:MAX_TOTAL_CHARS]
        warnings.append(
            f"⚠️ Combined intake exceeded {MAX_TOTAL_CHARS:,} chars and was truncated."
        )

    return combined, warnings
